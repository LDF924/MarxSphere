# xlsx2csv.py — 统一分析台 xlsx 读取(2026-09-09)
# 用法: python xlsx2csv.py <input.xlsx路径> → stdout CSV(首个含数据的 sheet)
# 依赖: openpyxl(实证 venv 已装)
import sys
import io

def main():
    path = sys.argv[1]
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("ERR:openpyxl 未安装", file=sys.stderr)
        sys.exit(2)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print(f"ERR:无法打开 xlsx: {e}", file=sys.stderr)
        sys.exit(3)
    ws = None
    for name in wb.sheetnames:
        cand = wb[name]
        if cand.max_row and cand.max_column:
            ws = cand
            break
    if ws is None:
        print("ERR:无数据 sheet", file=sys.stderr)
        sys.exit(4)
    out = io.StringIO()
    import csv
    w = csv.writer(out, lineterminator="\n")
    for row in ws.iter_rows(values_only=True):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        w.writerow(["" if v is None else v for v in row])
    print(out.getvalue(), end="")

if __name__ == "__main__":
    main()
